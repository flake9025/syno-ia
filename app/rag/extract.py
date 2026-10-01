"""Extraction du texte des formats bureautiques courants.

Chaque extracteur renvoie une liste de `TextSegment` afin de conserver la
localisation d'origine (numéro de page, nom d'onglet, numéro de diapositive),
qui sera ensuite affichée dans les citations de la réponse.
"""

from __future__ import annotations

import csv
import io
import json
import logging
from pathlib import PurePosixPath

from .chunker import TextSegment

logger = logging.getLogger(__name__)

PLAIN_TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".ini", ".cfg", ".conf", ".env",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".h", ".cpp", ".hpp", ".cs",
    ".go", ".rs", ".rb", ".php", ".sh", ".bash", ".ps1", ".sql", ".yaml", ".yml",
    ".toml", ".tex", ".srt", ".vtt",
}
SUPPORTED_EXTENSIONS = PLAIN_TEXT_EXTENSIONS | {
    ".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".csv", ".tsv", ".json", ".html", ".htm", ".xml",
}

#: Au-delà de cette taille, un fichier texte est tronqué (protection mémoire).
MAX_TEXT_CHARS = 2_000_000


class UnsupportedDocument(Exception):
    """Le format du fichier n'est pas pris en charge."""


def is_supported(name: str) -> bool:
    return PurePosixPath(name).suffix.lower() in SUPPORTED_EXTENSIONS


def extension_of(name: str) -> str:
    return PurePosixPath(name).suffix.lower()


def _decode(data: bytes) -> str:
    try:
        from charset_normalizer import from_bytes

        best = from_bytes(data).best()
        if best is not None:
            return str(best)
    except Exception:  # pragma: no cover - dépendance optionnelle
        pass
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# --------------------------------------------------------------- extracteurs
def _extract_pdf(data: bytes) -> list[TextSegment]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as exc:  # pragma: no cover - PDF protégé
            raise UnsupportedDocument("PDF protégé par mot de passe") from exc

    segments: list[TextSegment] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # pragma: no cover - page corrompue
            logger.debug("Page PDF %d illisible : %s", number, exc)
            continue
        if text.strip():
            segments.append(TextSegment(text=text, location=f"page {number}"))
    return segments


def _extract_docx(data: bytes) -> list[TextSegment]:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [TextSegment(text="\n".join(parts))] if parts else []


def _extract_pptx(data: bytes) -> list[TextSegment]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    segments: list[TextSegment] = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if parts:
            segments.append(TextSegment(text="\n".join(parts), location=f"diapositive {number}"))
    return segments


def _extract_xlsx(data: bytes) -> list[TextSegment]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    segments: list[TextSegment] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value not in (None, "")]
                if values:
                    rows.append(" | ".join(values))
                if len(rows) >= 5000:  # garde-fou mémoire
                    break
            if rows:
                segments.append(TextSegment(text="\n".join(rows), location=f"onglet {sheet.title}"))
    finally:
        workbook.close()
    return segments


def _extract_csv(data: bytes, delimiter: str | None = None) -> list[TextSegment]:
    text = _decode(data)[:MAX_TEXT_CHARS]
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=delimiter or ",;\t|")
        reader = csv.reader(io.StringIO(text), dialect)
    except csv.Error:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter or ",")
    rows = []
    for row in reader:
        values = [cell.strip() for cell in row if cell and cell.strip()]
        if values:
            rows.append(" | ".join(values))
    return [TextSegment(text="\n".join(rows))] if rows else []


def _extract_html(data: bytes) -> list[TextSegment]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data)[:MAX_TEXT_CHARS], "html.parser")
    for tag in soup(["script", "style", "noscript", "svg"]):
        tag.decompose()
    title = soup.title.string.strip() if soup.title and soup.title.string else ""
    body = soup.get_text("\n")
    text = f"{title}\n{body}" if title else body
    return [TextSegment(text=text)] if text.strip() else []


def _extract_json(data: bytes) -> list[TextSegment]:
    raw = _decode(data)[:MAX_TEXT_CHARS]
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return [TextSegment(text=raw)] if raw.strip() else []
    return [TextSegment(text=json.dumps(parsed, ensure_ascii=False, indent=1))]


def _extract_plain(data: bytes) -> list[TextSegment]:
    text = _decode(data)[:MAX_TEXT_CHARS]
    return [TextSegment(text=text)] if text.strip() else []


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".xlsx": _extract_xlsx,
    ".xlsm": _extract_xlsx,
    ".csv": _extract_csv,
    ".tsv": lambda data: _extract_csv(data, delimiter="\t"),
    ".json": _extract_json,
    ".html": _extract_html,
    ".htm": _extract_html,
    ".xml": _extract_html,
}


def extract(name: str, data: bytes) -> list[TextSegment]:
    """Extrait le texte d'un document à partir de son nom et de son contenu binaire."""
    extension = extension_of(name)
    extractor = _EXTRACTORS.get(extension)
    if extractor is None:
        if extension in PLAIN_TEXT_EXTENSIONS:
            return _extract_plain(data)
        raise UnsupportedDocument(f"Format non pris en charge : {extension or 'inconnu'}")
    try:
        return extractor(data)
    except UnsupportedDocument:
        raise
    except Exception as exc:
        raise UnsupportedDocument(f"Extraction impossible ({type(exc).__name__}: {exc})") from exc
